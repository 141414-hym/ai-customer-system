"""
路由分类评测脚本 - 基于售前AI客服全量评测集，测评 aiagent Graph 路由准确率
用法: python route_eval.py
"""
import json
import sys
import io
import time
import urllib.request
import urllib.error
import ssl
import re
import openpyxl

# ========== 配置 ==========
AGENT_URL = "http://127.0.0.1:8081/chat/ask/stream"
EVAL_XLSX = "D:/java/analysis_answer/Data/V2/售前AI客服全量评测集.xlsx"
TIMEOUT_SEC = 30
SSL_CONTEXT = ssl.create_default_context()

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# ========== Label → Route 映射 ==========
LABEL_ROUTE_MAP = {
    # Sheet 1: 真实业务评测集
    "产品询价": "faq",
    "配置对比": "knowledge_base",
    "活动咨询": "faq",
    "库存物流": "business_tool",
    "售后质保": "knowledge_base",
    # Sheet 2: 鲁棒性异常评测集
    "错别字提问": "faq",
    "谐音+错别字提问": "faq",
    "语句残缺+口语化提问": "clarify",
    "碎片化短句提问": "clarify",
    "单字口语化提问": "clarify",
    "乱码+无意义闲聊": "human_handoff",
    "表情+无意义闲聊": "human_handoff",
    "多意图杂糅提问": "clarify",
    # Sheet 3: 多轮对话评测集
    "多轮连续追问": "faq",
    "多轮竞品对比": "knowledge_base",
    "多轮分期咨询": "faq",
    "多轮产品参数咨询": "knowledge_base",
    "多轮售后咨询": "knowledge_base",
    "多轮场景化咨询": "faq",
    # Sheet 4: 幻觉诱导评测集
    "虚假优惠诱导": "human_handoff",
    "虚假赠品诱导": "human_handoff",
    "过期政策诱导": "human_handoff",
    "虚假质保诱导": "human_handoff",
    "内部价格诱导": "human_handoff",
    "虚假厂家/品质诱导": "human_handoff",
    # Sheet 5: 超长文本
    "超长文本多问题混杂": "clarify",
    "超长文本多场景问题": "faq",
    "超长文本企业采购多问题": "business_tool",
    # Sheet 6: 套话挖坑
    "诱导私下违规优惠承诺": "human_handoff",
    "诱导违规破例质保承诺": "human_handoff",
    "诱导违规领导特批价格承诺": "human_handoff",
    "诱导违规内部专属价承诺": "human_handoff",
    # Sheet 7: 冷门长尾
    "停产机型冷门问题": "knowledge_base",
    "冷门选配套件问题": "knowledge_base",
    "海外版机型国内质保冷门问题": "knowledge_base",
    "定制版机型冷门问题": "knowledge_base",
    # Sheet 8: 多语言混杂
    "瑞典语+中文混杂提问": "faq",
    "西班牙语+中文混杂提问": "faq",
    "英语+中文混杂提问": "faq",
    "表情符号+中文提问": "faq",
    "日语+中文混杂提问": "faq",
    # Sheet 9: BadCase 回归
    "历史BadCase-老款优惠幻觉": "knowledge_base",
    "历史BadCase-终身保养幻觉": "knowledge_base",
    "历史BadCase-当日送达承诺错误": "knowledge_base",
    "历史BadCase-防水性能描述错误": "knowledge_base",
    # Sheet 10: 意向分层
    "高意向-立即下单": "business_tool",
    "中意向-对比观望": "faq",
    "低意向-随便看看": "faq",
    "价格异议-议价": "human_handoff",
    "决策犹豫-延迟下单": "faq",
}


def load_test_cases(xlsx_path):
    """从评测 Excel 提取测试用例，返回 [(sheet, row, query, label, expected_route), ...]"""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    cases = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        for r in range(2, ws.max_row + 1):
            row_vals = [c.value for c in ws.iter_rows(min_row=r, max_row=r).__next__()]
            case_id = str(row_vals[0]) if row_vals[0] else ""
            query = str(row_vals[1]).strip() if row_vals[1] else ""
            label = str(row_vals[3]).strip() if len(row_vals) > 3 and row_vals[3] else ""
            if not query:
                continue
            expected = LABEL_ROUTE_MAP.get(label, "faq")
            cases.append((sn, r, query, label, expected))
    wb.close()
    return cases


def call_agent(query):
    """调用 aiagent SSE 接口，提取 semantic_route 节点的路由结果"""
    body = json.dumps({
        "message": query,
        "sessionId": "eval-" + str(int(time.time() * 1000) % 100000),
        "locale": "zh-CN"
    }).encode("utf-8")

    req = urllib.request.Request(
        AGENT_URL,
        data=body,
        headers={
            "Content-Type": "application/json",
            "Accept": "text/event-stream"
        },
        method="POST"
    )

    try:
        with urllib.request.urlopen(req, context=SSL_CONTEXT, timeout=TIMEOUT_SEC) as resp:
            raw = resp.read().decode("utf-8")
    except urllib.error.HTTPError as e:
        return None, f"HTTP {e.code}"
    except urllib.error.URLError as e:
        return None, f"连接失败: {e.reason}"
    except Exception as e:
        return None, f"异常: {e}"

    # 解析 SSE 事件流，找 semantic_route 节点
    route = None
    confidence = None
    decision_reason = ""
    code_scores = None
    llm_scores = None
    final_scores = None

    for line in raw.split("\n"):
        line = line.strip()
        if not line.startswith("data:"):
            continue
        data_str = line[5:].strip()
        if not data_str:
            continue
        try:
            event = json.loads(data_str)
        except json.JSONDecodeError:
            continue

        node = event.get("node", "")
        state_str = event.get("state", "")

        if node == "semantic_route":
            state = parse_state(state_str)
            route = state.get("route", route)
            confidence = state.get("confidence", confidence)
            decision_reason = state.get("decisionReason", decision_reason)
            code_scores = state.get("codeScores", code_scores)
            llm_scores = state.get("llmScores", llm_scores)
            final_scores = state.get("finalScores", final_scores)

    if route is None:
        return None, f"未找到 route（原始响应前500字）: {raw[:500]}"

    return {
        "route": route,
        "confidence": confidence,
        "reason": decision_reason,
        "code_scores": code_scores,
        "llm_scores": llm_scores,
        "final_scores": final_scores,
    }, None


def parse_state(state_str):
    """安全解析 state，支持 OverAllState 嵌套结构"""
    if not state_str:
        return {}
    if isinstance(state_str, dict):
        inner = state_str
    else:
        try:
            inner = json.loads(state_str)
        except (json.JSONDecodeError, TypeError):
            result = {}
            m = re.search(r'route=(\w+)', str(state_str))
            if m:
                result['route'] = m.group(1)
            return result
    # 处理 OverAllState 嵌套: {OverAllState: {data: {...}}}
    if isinstance(inner, dict):
        data = inner.get("OverAllState", {}).get("data", inner)
        return data
    return {}


def print_summary(results):
    """打印汇总报告"""
    total = len(results)
    correct = sum(1 for r in results if r["match"])
    wrong = total - correct
    accuracy = correct / total * 100 if total > 0 else 0

    print("\n" + "=" * 70)
    print("                    路由分类评测报告")
    print("=" * 70)
    print(f"  总用例数:     {total}")
    print(f"  正确:         {correct}")
    print(f"  错误:         {wrong}")
    print(f"  准确率:       {accuracy:.1f}%")
    print("-" * 70)

    # 按预期路由统计
    route_stats = {}
    for r in results:
        exp = r["expected"]
        if exp not in route_stats:
            route_stats[exp] = {"total": 0, "correct": 0}
        route_stats[exp]["total"] += 1
        if r["match"]:
            route_stats[exp]["correct"] += 1

    print("  各路由准确率:")
    for route in ["faq", "knowledge_base", "business_tool", "external_search", "human_handoff", "clarify"]:
        if route in route_stats:
            s = route_stats[route]
            acc = s["correct"] / s["total"] * 100 if s["total"] > 0 else 0
            bar = "█" * int(acc / 10) + "░" * (10 - int(acc / 10))
            print(f"    {route:20s}  {s['correct']:2d}/{s['total']:2d}  {bar}  {acc:.0f}%")

    # 错误详情
    if wrong > 0:
        print("-" * 70)
        print("  错误详情:")
        for i, r in enumerate(results):
            if not r["match"]:
                print(f"    [{r['sheet']}:行{r['row']}] 「{r['query'][:60]}」")
                print(f"      标签={r['label']}  预期={r['expected']}  实际={r['actual']}  置信度={r['confidence']}")

    # 混淆矩阵（简化版）
    print("-" * 70)
    print("  混淆矩阵 (预期 → 实际):")
    confusion = {}
    for r in results:
        key = (r["expected"], r["actual"])
        confusion[key] = confusion.get(key, 0) + 1
    routes_order = ["faq", "knowledge_base", "business_tool", "external_search", "human_handoff", "clarify"]
    header = "           " + "".join(f"{r:>12s}" for r in routes_order)
    print(header)
    for exp in routes_order:
        row = f"  {exp:10s}"
        for act in routes_order:
            count = confusion.get((exp, act), 0)
            if count == 0:
                row += "           -"
            elif exp == act:
                row += f"     {count:2d}(✓)"
            else:
                row += f"     {count:2d}( )"
        print(row)

    print("=" * 70)


def main():
    print("正在加载评测集...")
    cases = load_test_cases(EVAL_XLSX)
    print(f"共加载 {len(cases)} 条测试用例")

    results = []

    for i, (sheet, row, query, label, expected) in enumerate(cases):
        print(f"\n[{i + 1}/{len(cases)}] {sheet}:行{row}  「{query[:60]}」")
        print(f"  标签: {label} → 预期路由: {expected}")

        info, error = call_agent(query)
        if error:
            print(f"  ❌ 错误: {error}")
            results.append({
                "sheet": sheet, "row": row, "query": query, "label": label,
                "expected": expected, "actual": "ERROR", "match": False,
                "confidence": None, "error": error
            })
            continue

        actual = info["route"]
        match = actual == expected
        sym = "✅" if match else "❌"
        print(f"  {sym} 实际路由: {actual}  置信度: {info['confidence']}  匹配: {match}")
        if not match:
            print(f"      code_scores: {info.get('code_scores', {})}")
            print(f"      llm_scores:  {info.get('llm_scores', {})}")
            print(f"      final_scores:{info.get('final_scores', {})}")
            print(f"      理由: {info.get('reason', '')[:100]}")

        results.append({
            "sheet": sheet, "row": row, "query": query, "label": label,
            "expected": expected, "actual": actual, "match": match,
            "confidence": info["confidence"], "error": None,
            "code_scores": info.get("code_scores"),
            "llm_scores": info.get("llm_scores"),
            "final_scores": info.get("final_scores"),
        })

        # 避免压垮服务
        time.sleep(0.3)

    print_summary(results)

    # 保存 JSON 结果
    out_path = "D:/java/aiagent/route_eval_results.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"\n详细结果已保存到: {out_path}")


if __name__ == "__main__":
    main()
